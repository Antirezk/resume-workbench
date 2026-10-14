import threading
from pathlib import Path

from PIL import Image, ImageOps
from pypdf import PdfReader


_ocr = None
_ocr_lock = threading.Lock()
_pdf_lock = threading.Lock()
SUPPORTED = {".pdf", ".docx", ".txt", ".png", ".jpg", ".jpeg"}


def render_pdf_page(path, index):
    # PDFium is process-global and must not be called concurrently from workers.
    import pypdfium2
    with _pdf_lock:
        pdf = pypdfium2.PdfDocument(str(path))
        page, bitmap = None, None
        try:
            page = pdf[index]
            width, height = page.get_size()
            bitmap = page.render(scale=min(2.2, 2800 / max(width, height)))
            return bitmap.to_pil().copy()
        finally:
            if bitmap is not None:
                bitmap.close()
            if page is not None:
                page.close()
            pdf.close()


def ocr_image(image):
    global _ocr
    import numpy as np
    from rapidocr import RapidOCR
    image = ImageOps.exif_transpose(image).convert("RGB")
    image.thumbnail((2800, 2800))
    with _ocr_lock:
        if _ocr is None:
            _ocr = RapidOCR()
        output = _ocr(np.asarray(image))
        texts = getattr(output, "txts", None)
    return "\n".join(texts) if texts is not None else ""


def extract_text(path, use_ocr=True, force_ocr=False):
    path = Path(path)
    suffix = path.suffix.lower()
    notes, pages, method = [], 1, "文字提取"
    if suffix == ".pdf":
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("PDF 已加密，请先移除密码")
        pages = len(reader.pages)
        if pages > 30:
            raise ValueError("单份简历最多支持 30 页")
        parts = []
        for index, page in enumerate(reader.pages):
            try:
                text = page.extract_text(extraction_mode="layout") or ""
            except Exception:
                text = page.extract_text() or ""
            text_chars = sum(c.isalnum() for c in text)
            if force_ocr or text_chars < 30:
                if not use_ocr:
                    raise ValueError(f"第 {index + 1} 页缺少文字，请启用扫描件识别")
                image = render_pdf_page(path, index)
                try:
                    text = ocr_image(image)
                finally:
                    image.close()
                method = "OCR" if force_ocr else "文字提取 + OCR"
                if sum(c.isalnum() for c in text) < 10:
                    raise ValueError(f"第 {index + 1} 页无法识别，请检查扫描清晰度或移除空白页")
            parts.append(f"[第 {index + 1} 页]\n{text}")
        text = "\n\n".join(parts)
    elif suffix == ".docx":
        from docx import Document
        doc = Document(path)
        text = "\n".join([p.text for p in doc.paragraphs] +
                         [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
    elif suffix == ".txt":
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "gb18030", "utf-16"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeError:
                continue
        else:
            raise ValueError("无法读取文本编码，请保存为 UTF-8")
    elif suffix in {".png", ".jpg", ".jpeg"}:
        if not use_ocr:
            raise ValueError("图片简历需要启用扫描件识别")
        with Image.open(path) as image:
            text = ocr_image(image)
        method = "OCR"
    else:
        raise ValueError("不支持的文件格式，请使用 PDF、DOCX、TXT、JPG 或 PNG")
    text = text.replace("\x00", "").strip()
    if sum(c.isalnum() for c in text) < 20:
        raise ValueError("提取的文字过少，请检查文件或启用强制 OCR")
    if len(text) > 60000:
        raise ValueError("简历文字超过 60,000 字符，请拆分文档后上传")
    return {"text": text, "method": method, "pages": pages, "notes": notes}
